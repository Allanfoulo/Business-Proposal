import io
import os
import re
import docx
from fpdf import FPDF
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def strip_md(text: str) -> str:
    text = text.replace("**", "")
    text = text.replace("*", "")
    text = text.replace("#", "")
    text = re.sub(r'([!*_=~-])', r'\\\1', text)
    return text


def process_markdown_to_runs(paragraph, text: str) -> None:
    text = text.replace("\\", "")
    bold_pattern = re.compile(r'\*\*(.*?)\*\*')
    last_end = 0

    for match in bold_pattern.finditer(text):
        if match.start() > last_end:
            paragraph.add_run(text[last_end:match.start()])
        paragraph.add_run(match.group(1)).bold = True
        last_end = match.end()

    if last_end < len(text):
        paragraph.add_run(text[last_end:])


def add_markdown_table(doc, text: str) -> None:
    rows = [row.strip() for row in text.split('\n') if row.strip() and '|' in row]
    if not rows:
        return

    # Skip markdown header separator rows like `---|---` or `:---|---:`
    def _is_separator_row(r: str) -> bool:
        cells = [c.strip() for c in r.split('|')]
        return all(re.match(r'^:?-{3,}:?$', c) for c in cells if c)

    filtered_rows = [r for r in rows if not _is_separator_row(r)]
    if not filtered_rows:
        return

    table_data = [[cell.strip() for cell in row.split('|') if cell.strip()] for row in filtered_rows]
    if not table_data:
        return

    num_rows = len(table_data)
    num_cols = len(table_data[0])
    table = doc.add_table(rows=num_rows, cols=num_cols)

    for row_idx, row_data in enumerate(table_data):
        for col_idx, cell_data in enumerate(row_data[:num_cols]):
            table.cell(row_idx, col_idx).text = cell_data


# Helper: apply a basic Word theme with margins and header
def apply_docx_theme(doc):
    try:
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
    except Exception:
        pass

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        header = section.header
        if header.paragraphs:
            header.paragraphs[0].text = "Business Proposal"
            header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


# Helper: inline markdown (bold, italic, code, links) → docx runs
def add_runs_from_inline_markdown(paragraph, text: str) -> None:
    text = text.replace("\\", "")
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))")
    last = 0
    for m in pattern.finditer(text):
        if m.start() > last:
            paragraph.add_run(text[last:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*"):
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        elif token.startswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
        elif token.startswith("["):
            mt = re.match(r"\[(.*?)\]\((.*?)\)", token)
            if mt:
                paragraph.add_run(mt.group(1))
                paragraph.add_run(f" ({mt.group(2)})").italic = True
        last = m.end()
    if last < len(text):
        paragraph.add_run(text[last:])


# Helper: render markdown blocks into DOCX with headings, lists, tables, code
def render_markdown_to_docx(doc, text: str) -> None:
    lines = text.splitlines()
    in_code = False
    code_lines = []
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_markdown_table(doc, "\n".join(table_buf))
            table_buf = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_lines))
                r.font.name = "Consolas"
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue

        if "|" in line:
            table_buf.append(line)
            continue
        else:
            flush_table()

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), min(level, 9))
            continue

        if re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="ListBullet")
            continue

        if re.match(r"^\s*\d+\.\s+", line):
            doc.add_paragraph(re.sub(r"^\s*\d+\.\s+", "", line), style="ListNumber")
            continue

        p = doc.add_paragraph()
        add_runs_from_inline_markdown(p, line)

    flush_table()


class StyledPDF(FPDF):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._font_name = "helvetica"
        self._has_bold = False
        self.title_text = "Business Proposal"

    def set_branding(self, font_name: str, has_bold: bool, title: str) -> None:
        self._font_name = font_name
        self._has_bold = has_bold
        self.title_text = title

    def header(self):
        self.set_y(12)
        if self._has_bold:
            self.set_font(self._font_name, "B", 12)
        else:
            self.set_font(self._font_name, size=12)
        self.set_text_color(60, 60, 60)
        self.cell(0, 8, self.title_text, 0, 1, "C")
        self.ln(2)
        self.set_draw_color(200, 200, 200)
        self.line(10, 24, self.w - 10, 24)

    def footer(self):
        self.set_y(-15)
        self.set_font(self._font_name, size=10)
        self.set_text_color(120, 120, 120)
        self.cell(0, 10, f"Page {self.page_no()}", 0, 0, "C")


def _render_table_to_pdf(pdf: FPDF, table_text: str, font_name: str) -> None:
    rows = [r for r in table_text.split("\n") if "|" in r]
    if not rows:
        return
    table = [[c.strip() for c in r.split("|") if c.strip()] for r in rows]
    col_count = max(len(r) for r in table)
    available_width = pdf.w - pdf.l_margin - pdf.r_margin
    col_width = available_width / max(col_count, 1)
    line_height = 7

    pdf.set_font(font_name, size=11)
    for i, row in enumerate(table):
        if i == 0:
            pdf.set_fill_color(230, 230, 230)
        else:
            pdf.set_fill_color(250, 250, 250) if i % 2 == 0 else pdf.set_fill_color(240, 240, 240)
        for j in range(col_count):
            cell_text = row[j] if j < len(row) else ""
            pdf.cell(col_width, line_height, _sanitize_if_needed(cell_text, unicode_enabled=(font_name != "helvetica")), border=1, align="L", fill=True)
        pdf.ln(line_height)
    pdf.ln(4)


def render_markdown_to_pdf(pdf: FPDF, text: str, font_name: str, has_bold: bool) -> None:
    lines = text.splitlines()
    in_code = False
    code_lines = []
    table_buf = []
    bullet_char = "•" if font_name != "helvetica" else "*"

    def flush_table():
        nonlocal table_buf
        if table_buf:
            _render_table_to_pdf(pdf, "\n".join(table_buf), font_name)
            table_buf = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                pdf.set_font(font_name, size=11)
                pdf.set_text_color(50, 50, 50)
                for cl in code_lines:
                    pdf.set_x(pdf.l_margin + 4)
                    pdf.multi_cell(0, 6, _sanitize_if_needed(cl, unicode_enabled=(font_name != "helvetica")))
                pdf.set_text_color(0, 0, 0)
                code_lines = []
                in_code = False
                pdf.ln(2)
            else:
                flush_table()
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue

        if "|" in line:
            table_buf.append(line)
            continue
        else:
            flush_table()

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            size = {1: 16, 2: 14, 3: 13}.get(level, 12)
            if has_bold:
                pdf.set_font(font_name, "B", size)
            else:
                pdf.set_font(font_name, size=size)
            pdf.set_text_color(30, 30, 30)
            pdf.multi_cell(0, 8, _sanitize_if_needed(m.group(2).strip(), unicode_enabled=(font_name != "helvetica")))
            pdf.ln(2)
            pdf.set_font(font_name, size=12)
            pdf.set_text_color(0, 0, 0)
            continue

        if re.match(r"^\s*[-*]\s+", line):
            pdf.set_x(pdf.l_margin + 4)
            pdf.cell(5, 6, bullet_char, 0, 0)
            txt = re.sub(r"^\s*[-*]\s+", "", line)
            pdf.multi_cell(0, 6, _sanitize_if_needed(txt, unicode_enabled=(font_name != "helvetica")))
            continue

        if re.match(r"^\s*\d+\.\s+", line):
            txt = re.sub(r"^\s*\d+\.\s+", "", line)
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(0, 6, _sanitize_if_needed(f"- {txt}", unicode_enabled=(font_name != "helvetica")))
            continue

        pdf.multi_cell(0, 6, _sanitize_if_needed(line, unicode_enabled=(font_name != "helvetica")))
        pdf.ln(1)

    flush_table()


def build_pdf(section_map, section_order, section_to_header) -> bytes:
    pdf = StyledPDF()
    pdf.set_auto_page_break(True, margin=15)
    pdf.set_margins(15, 20, 15)

    # Prefer Unicode-capable TrueType fonts to avoid encoding errors
    font_name, has_bold = _setup_unicode_font(pdf)
    pdf.set_branding(font_name, has_bold, "Business Proposal")
    pdf.alias_nb_pages()
    pdf.add_page()

    for key in section_order:
        header_text = section_to_header.get(key, "Section")
        content_text = section_map.get(key, f"[Missing: {key}]")

        # Header
        if has_bold:
            pdf.set_font(font_name, 'B', 14)
        else:
            pdf.set_font(font_name, size=14)
        header_safe = _sanitize_if_needed(header_text, unicode_enabled=(font_name != "helvetica"))
        pdf.set_text_color(25, 25, 25)
        pdf.multi_cell(0, 10, header_safe)
        pdf.set_text_color(0, 0, 0)

        # Body
        pdf.set_font(font_name, size=12)
        render_markdown_to_pdf(pdf, content_text, font_name, has_bold)
        pdf.ln(4)

    pdf_out = pdf.output(dest='S')
    if isinstance(pdf_out, (bytes, bytearray)):
        return bytes(pdf_out)
    return str(pdf_out).encode('latin1')


def build_docx(section_map, section_order, section_to_header) -> bytes:
    doc = docx.Document()
    doc.add_heading("Business Proposal", 0)
    apply_docx_theme(doc)

    for key in section_order:
        header_text = section_to_header.get(key, "Section")
        doc.add_heading(header_text, 3)
        section_text = section_map.get(key, f"[Missing: {key}]")
        render_markdown_to_docx(doc, section_text)

    mem_file = io.BytesIO()
    doc.save(mem_file)
    mem_file.seek(0)
    return mem_file.getvalue()


def _setup_unicode_font(pdf: FPDF):
    """Register a Unicode-capable font if available.

    Tries Windows Arial TTFs first; falls back to core 'helvetica' if not found.
    Returns (font_name, has_bold_variant).
    """
    # Common Windows font locations
    arial = r"C:\\Windows\\Fonts\\arial.ttf"
    arial_bold = r"C:\\Windows\\Fonts\\arialbd.ttf"

    has_bold = False
    try:
        if os.path.exists(arial):
            pdf.add_font("ArialUnicode", "", arial, uni=True)
            if os.path.exists(arial_bold):
                pdf.add_font("ArialUnicode", "B", arial_bold, uni=True)
                has_bold = True
            return "ArialUnicode", has_bold
    except Exception:
        # Ignore font registration errors and fall back
        pass

    # Fallback: core font (latin-1 only)
    return "helvetica", False


def _sanitize_if_needed(text: str, unicode_enabled: bool) -> str:
    """Ensure text is safe for PDF when Unicode is not enabled.

    If unicode_enabled is False, convert to latin-1 safely and replace common
    problematic characters.
    """
    if unicode_enabled:
        return text

    # Replace some common Unicode punctuation with ASCII equivalents
    replacements = {
        "–": "-",
        "—": "-",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "…": "...",
        "•": "*",
        "✓": "check",
        "→": "->",
        "←": "<-",
        "×": "x",
    }
    for k, v in replacements.items():
        text = text.replace(k, v)

    # Drop any remaining non-latin-1 characters
    return text.encode('latin-1', 'ignore').decode('latin-1')